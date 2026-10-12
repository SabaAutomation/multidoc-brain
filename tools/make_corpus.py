"""Generate the synthetic test corpus: a fictional company, "Cedar & Co. Outfitters".

Everything here is invented. The answers are known, so scoring needs no judge.
Run:  python tools/make_corpus.py
"""
import csv
import json
import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "corpus"
OUT.mkdir(exist_ok=True)


def make_docx_handbook():
    from docx import Document
    d = Document()
    d.add_heading("Cedar & Co. Outfitters - Employee Handbook", 1)
    d.add_paragraph("Standard working hours are 9:00 to 17:30, Monday to Friday.")
    d.add_paragraph("New employees complete a probation period of 90 days.")
    d.add_paragraph("Employees may work remotely on Wednesdays and Fridays with manager approval.")
    d.add_paragraph("Leave entitlements")
    t = d.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, ["Role", "Annual leave days", "Sick days"]):
        c.text = h
    for role, a, s in [("Associate", "18", "8"), ("Senior", "22", "10"), ("Manager", "26", "12")]:
        cells = t.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = role, a, s
    d.add_paragraph("Unused leave does not carry over to the next calendar year.")
    d.save(OUT / "employee_handbook.docx")


def make_docx_meeting():
    from docx import Document
    d = Document()
    d.add_heading("Leadership meeting notes - September 2025", 1)
    d.add_paragraph("The board approved opening a new store in Denver. The opening date is March 3.")
    d.add_paragraph("The approved budget for the Denver store is $240,000.")
    d.add_paragraph("Priya Nair is the owner of the Denver opening.")
    d.add_paragraph("The team will hire 6 associates for the new store.")
    d.save(OUT / "meeting_notes_2025-09.docx")


def make_pdf_expense():
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(OUT / "expense_policy.pdf"), pagesize=A4)
    y = 800
    lines = [
        "Expense Policy",
        "",
        "Meals are reimbursed up to $45 per day while travelling.",
        "Flights must be economy class unless the flight is longer than 6 hours.",
        "Receipts are required for every expense above $25.",
        "",
        "Approvals",
        "A manager approves expenses up to $500.",
        "A director must approve any expense above $500.",
        "Expense claims must be submitted within 30 days of the purchase.",
    ]
    for ln in lines:
        c.setFont("Helvetica-Bold" if ln in ("Expense Policy", "Approvals") else "Helvetica", 12)
        c.drawString(60, y, ln)
        y -= 22
    c.save()


def make_xlsx_prices():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(["SKU", "Product", "Price", "Stock"])
    for r in [("CJ-101", "Trail Jacket", 129, 42), ("SB-220", "Summit Backpack", 89.5, 17),
              ("RB-310", "Ridge Boots", 154, 8), ("CS-045", "Camp Stove", 64.99, 0),
              ("IB-009", "Insulated Bottle", 24, 120)]:
        ws.append(list(r))
    w2 = wb.create_sheet("Discounts")
    w2.append(["Tier", "Minimum order", "Discount"])
    for r in [("Bronze", 500, "5%"), ("Silver", 1500, "8%"), ("Gold", 5000, "12%")]:
        w2.append(list(r))
    wb.save(OUT / "price_list.xlsx")


def make_xlsx_sales():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Q3"
    ws.append(["Region", "Revenue", "Units"])
    for r in [("North", 182400, 1310), ("South", 141900, 1022), ("West", 205750, 1498)]:
        ws.append(list(r))
    wb.save(OUT / "quarterly_sales.xlsx")


def make_csv_vendors():
    with open(OUT / "vendors.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Vendor", "Category", "Lead time days", "Contact"])
        w.writerows([
            ["Alpine Textiles", "fabric", 21, "orders@alpine-textiles.example"],
            ["Ridgeline Hardware", "hardware", 14, "sales@ridgeline.example"],
            ["Summit Packaging", "packaging", 9, "hello@summitpack.example"],
            ["Northwind Freight", "shipping", 5, "dispatch@northwind-freight.example"],
        ])


def make_md_faq():
    (OUT / "customer_faq.md").write_text(
        "# Customer FAQ\n\n"
        "Items can be returned within 30 days with the receipt.\n\n"
        "Orders over $75 qualify for free shipping.\n\n"
        "Customer support is open 08:00 to 20:00, Monday to Saturday.\n\n"
        "Ridge Boots carry a warranty of 2 years against manufacturing defects.\n",
        encoding="utf-8")


def make_scanned_pdf():
    """An image-only PDF (no text layer), like a photographed or scanned notice."""
    from PIL import Image, ImageDraw, ImageFont
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    img = Image.new("RGB", (1240, 700), "white")
    dr = ImageDraw.Draw(img)
    font = None
    for p in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "C:/Windows/Fonts/arial.ttf",
              "/System/Library/Fonts/Supplemental/Arial.ttf"):
        if os.path.exists(p):
            font = ImageFont.truetype(p, 44)
            break
    if font is None:
        font = ImageFont.load_default(size=44)
    lines = ["SAFETY NOTICE", "", "Fire assembly point: North car park.",
             "Evacuation drills are held every 6 months.", "Safety officer: Daniel Okoye."]
    y = 60
    for ln in lines:
        dr.text((60, y), ln, fill="black", font=font)
        y += 90
    tmp = OUT / "_scan.png"
    img.save(tmp)
    c = canvas.Canvas(str(OUT / "scanned_safety_notice.pdf"), pagesize=A4)
    c.drawImage(str(tmp), 20, 300, width=555, height=313)
    c.save()
    tmp.unlink()


QUESTIONS = [
    # prose (pdf / docx paragraphs / markdown)
    ("q01", "What are the standard working hours?", ["9:00", "17:30"], "prose", "employee_handbook.docx"),
    ("q02", "How long is the probation period?", ["90"], "prose", "employee_handbook.docx"),
    ("q03", "On which days can employees work remotely?", ["wednesday", "friday"], "prose", "employee_handbook.docx"),
    ("q04", "What is the daily meal reimbursement limit?", ["45"], "prose", "expense_policy.pdf"),
    ("q05", "Above what amount are receipts required?", ["25"], "prose", "expense_policy.pdf"),
    ("q06", "Who must approve expenses above $500?", ["director"], "prose", "expense_policy.pdf"),
    ("q07", "Within how many days must expense claims be submitted?", ["30"], "prose", "expense_policy.pdf"),
    ("q08", "How many days do customers have to return an item?", ["30"], "prose", "customer_faq.md"),
    ("q09", "What order value qualifies for free shipping?", ["75"], "prose", "customer_faq.md"),
    ("q10", "What are the customer support hours?", ["08:00|8:00|8 am|8am", "20:00|8 pm|8pm"], "prose", "customer_faq.md"),
    ("q11", "How long is the warranty on Ridge Boots?", ["2 years|two years|24 months"], "prose", "customer_faq.md"),
    ("q12", "On what date will the Denver store open?", ["march 3|3 march|3rd march|march 3rd"], "prose", "meeting_notes_2025-09.docx"),
    ("q13", "What is the approved budget for the Denver store?", ["240000"], "prose", "meeting_notes_2025-09.docx"),
    ("q14", "Who owns the Denver opening?", ["priya nair"], "prose", "meeting_notes_2025-09.docx"),
    ("q15", "How many associates will be hired for the new store?", ["6|six"], "prose", "meeting_notes_2025-09.docx"),
    # table inside a docx
    ("q16", "How many annual leave days does a Senior employee get?", ["22"], "docx_table", "employee_handbook.docx"),
    ("q17", "How many sick days does a Manager get?", ["12"], "docx_table", "employee_handbook.docx"),
    ("q18", "How many annual leave days does an Associate get?", ["18"], "docx_table", "employee_handbook.docx"),
    # spreadsheets
    ("q19", "What is the price of the Ridge Boots?", ["154"], "xlsx", "price_list.xlsx"),
    ("q20", "How many Summit Backpacks are in stock?", ["17"], "xlsx", "price_list.xlsx"),
    ("q21", "Which product is out of stock?", ["camp stove"], "xlsx", "price_list.xlsx"),
    ("q22", "What discount does the Silver tier get?", ["8%|8 percent|8 %"], "xlsx", "price_list.xlsx"),
    ("q23", "What is the minimum order for the Gold tier?", ["5000"], "xlsx", "price_list.xlsx"),
    ("q24", "What was the West region's revenue in Q3?", ["205750"], "xlsx", "quarterly_sales.xlsx"),
    ("q25", "How many units did the South region sell in Q3?", ["1022"], "xlsx", "quarterly_sales.xlsx"),
    # csv
    ("q26", "What is the lead time for Summit Packaging?", ["9"], "csv", "vendors.csv"),
    ("q27", "Which vendor supplies fabric?", ["alpine textiles"], "csv", "vendors.csv"),
    ("q28", "What is the lead time for Northwind Freight?", ["5"], "csv", "vendors.csv"),
    # scanned (image only)
    ("q29", "Where is the fire assembly point?", ["north car park"], "scanned", "scanned_safety_notice.pdf"),
    ("q30", "Who is the safety officer?", ["daniel okoye"], "scanned", "scanned_safety_notice.pdf"),
    ("q31", "How often are evacuation drills held?", ["6 months|six months"], "scanned", "scanned_safety_notice.pdf"),
    # answer is not in any document: the right reply is NOT FOUND
    ("q32", "What is the CEO's annual salary?", [], "unanswerable", ""),
    ("q33", "What is the price of the Alpine Tent?", [], "unanswerable", ""),
    ("q34", "When will the Seattle store open?", [], "unanswerable", ""),
]
GROUPS = ["prose", "docx_table", "xlsx", "csv", "scanned", "unanswerable"]


def write_questions():
    data = [{"id": i, "question": q, "gold": g, "kind": k, "file": f} for i, q, g, k, f in QUESTIONS]
    (OUT / "questions.json").write_text(json.dumps(data, indent=2), encoding="utf-8")


def main():
    make_docx_handbook()
    make_docx_meeting()
    make_pdf_expense()
    make_xlsx_prices()
    make_xlsx_sales()
    make_csv_vendors()
    make_md_faq()
    make_scanned_pdf()
    write_questions()
    print(f"Corpus written to {OUT} ({len(QUESTIONS)} questions)")


if __name__ == "__main__":
    sys.exit(main())
